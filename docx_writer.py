import os
from docx import Document
from config import Config
from utils import log_error, sanitize_filename, redact_secrets

def create_docx(transcript: str, original_file_name: str) -> str:
    """
    Creates a DOCX file containing the transcript.
    Returns the path to the generated DOCX file.
    """
    os.makedirs(Config.OUTPUT_DIR, exist_ok=True)
    
    clean_original_name = sanitize_filename(original_file_name)
    base_name = os.path.splitext(clean_original_name)[0]
    docx_filename = f"{base_name}.docx"
    docx_path = os.path.join(Config.OUTPUT_DIR, docx_filename)
    
    try:
        doc = Document()
        doc.add_heading(f"Transcript: {clean_original_name}", level=1)
        doc.add_paragraph(transcript)
        doc.save(docx_path)
        return docx_path
    except Exception as e:
        log_error("create_docx", redact_secrets(str(e)))
        raise
